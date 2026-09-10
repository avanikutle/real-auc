"""Step 7 — Spreadsheet export to Excel workbook.

Generates:
  - Master sheet: one row per property with key fields and calculated metrics
  - Per-property tabs: all extracted fields + Sources table with hyperlinks

Styling:
  - Font: Arial
  - Header: bold white text on navy background (#003366)
  - Wrapped text, thin borders, frozen header row
  - Hyperlinks to all source documents/URLs

Resilience:
  - Per-property tab failures are caught individually
  - Master sheet continues even if some properties fail
  - Missing-from-Master properties are logged
"""
from __future__ import annotations

import csv
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from auction_pipeline.db.models import LoanEstimate, Property, StepResult
from auction_pipeline.db.session import get_session

log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------
_NAVY = "003366"
_WHITE = "FFFFFF"
_LIGHT_BLUE = "DCE6F1"

_HEADER_FONT = Font(name="Arial", bold=True, color=_WHITE, size=10)
_BODY_FONT = Font(name="Arial", size=9)
_LINK_FONT = Font(name="Arial", size=9, color="0563C1", underline="single")

_HEADER_FILL = PatternFill("solid", fgColor=_NAVY)
_ALT_FILL = PatternFill("solid", fgColor=_LIGHT_BLUE)

_THIN = Side(style="thin")
_THIN_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="top")

# ---------------------------------------------------------------------------
# Master sheet column definitions
# ---------------------------------------------------------------------------
_MASTER_COLS = [
    ("Entry No", "entry_no"),
    ("Owner", "owner_full_name"),
    ("File No", "source_file_name"),
    ("R Number", "r_number"),
    ("Instrument No", "instrument_number"),
    ("Address", "address"),
    ("Legal Description", "legal_description"),
    ("County Assessed Value", "_assessed_value"),
    ("Trustee", "trustee"),
    ("Auction Price", "_auction_price"),        # placeholder — populated post-sale
    ("Approx Current Balance", "_est_balance"),
    ("Est % Paid Down", "_est_pct_paid_down"),
    ("Filter Flag", "_filter_flag"),
    ("Other Liens (count)", "_lien_count"),
    ("Sale Date", "sale_date"),
    ("Loan Type", "loan_type"),
    ("Lender", "lender"),
    ("Servicer", "servicer"),
    ("Is Purchase Money", "is_purchase_money"),
    ("Has HOA Rider", "has_hoa_rider"),
    ("Original Loan Amount", "original_loan_amount"),
    ("Origination Date", "loan_origination_date"),
    ("Assumed Rate", "_assumed_rate"),
    ("Monthly Payment (Est)", "_est_monthly_payment"),
    ("Status", "status"),
]


def _style_header(ws, row: int) -> None:
    """Apply header styling to the given row."""
    for cell in ws[row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER


def _set_column_widths(ws, widths: dict[str, int]) -> None:
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width


def _get_step_result(session, entry_no: str, step_name: str) -> dict:
    """Fetch extracted_json from a step_results row, or return {}."""
    sr = (
        session.query(StepResult)
        .filter_by(entry_no=entry_no, step_name=step_name)
        .first()
    )
    if sr and sr.extracted_json:
        try:
            return json.loads(sr.extracted_json)
        except (json.JSONDecodeError, TypeError):
            pass
    return {}


def _build_property_data(session, prop: Property) -> dict[str, Any]:
    """Collect all data for one property into a flat dict."""
    data: dict[str, Any] = {}

    # Basic property fields
    for col_key in ["entry_no", "owner_full_name", "source_file_name", "r_number",
                    "instrument_number", "address", "legal_description", "trustee",
                    "sale_date", "loan_type", "lender", "servicer", "is_purchase_money",
                    "has_hoa_rider", "original_loan_amount", "loan_origination_date", "status"]:
        data[col_key] = getattr(prop, col_key, None)

    # Loan estimate
    est: LoanEstimate | None = session.get(LoanEstimate, prop.entry_no)
    if est:
        data["_est_balance"] = est.est_remaining_balance
        data["_est_pct_paid_down"] = est.est_pct_paid_down
        data["_filter_flag"] = est.filter_flag
        data["_assumed_rate"] = est.assumed_rate
        data["_est_monthly_payment"] = est.est_monthly_payment
    else:
        data["_est_balance"] = None
        data["_est_pct_paid_down"] = None
        data["_filter_flag"] = "NOT_CALCULATED"
        data["_assumed_rate"] = None
        data["_est_monthly_payment"] = None

    # WCAD assessed value
    wcad = _get_step_result(session, prop.entry_no, "step2_wcad")
    data["_assessed_value"] = wcad.get("assessed_value")

    # Lien count
    liens_data = _get_step_result(session, prop.entry_no, "step5_liens")
    lien_count = liens_data.get("lien_count")
    data["_lien_count"] = lien_count

    # Auction price — not available pre-sale
    data["_auction_price"] = None

    # Collect all sources (URLs + file paths) from step_results
    sources = []
    for sr in prop.step_results:
        if sr.source_url:
            sources.append({"step": sr.step_name, "type": "url", "value": sr.source_url})
        if sr.source_file_path:
            sources.append({"step": sr.step_name, "type": "file", "value": sr.source_file_path})
    data["_sources"] = sources

    # All step_results JSON for detail tab
    data["_step_results"] = {
        sr.step_name: json.loads(sr.extracted_json or "{}") for sr in prop.step_results
    }

    return data


def _write_master_sheet(ws, properties_data: list[dict]) -> None:
    """Write the Master sheet."""
    ws.title = "Master"

    # Headers
    headers = [col[0] for col in _MASTER_COLS]
    for c, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=header)
    _style_header(ws, 1)
    ws.freeze_panes = "A2"

    # Data rows
    for r, data in enumerate(properties_data, start=2):
        fill = _ALT_FILL if r % 2 == 0 else None
        for c, (_, key) in enumerate(_MASTER_COLS, start=1):
            value = data.get(key)

            # Format specific types
            if key == "_est_pct_paid_down" and value is not None:
                value = round(value * 100, 2)  # as percent
            elif isinstance(value, bool):
                value = "Yes" if value else "No"

            cell = ws.cell(row=r, column=c, value=value)
            cell.font = _BODY_FONT
            cell.border = _THIN_BORDER
            cell.alignment = _WRAP
            if fill:
                cell.fill = fill

    # Column widths
    col_widths = {
        get_column_letter(1): 30,   # Entry No
        get_column_letter(2): 35,   # Owner
        get_column_letter(3): 22,   # File No
        get_column_letter(4): 12,   # R Number
        get_column_letter(5): 16,   # Instrument No
        get_column_letter(6): 35,   # Address
        get_column_letter(7): 50,   # Legal Description
        get_column_letter(8): 20,   # Assessed Value
        get_column_letter(9): 30,   # Trustee
        get_column_letter(10): 15,  # Auction Price
        get_column_letter(11): 22,  # Current Balance
        get_column_letter(12): 16,  # % Paid Down
        get_column_letter(13): 18,  # Filter Flag
        get_column_letter(14): 16,  # Lien Count
    }
    _set_column_widths(ws, col_widths)


def _write_property_sheet(wb: Workbook, data: dict) -> None:
    """Write a per-property detail tab."""
    entry_no = data.get("entry_no", "unknown")
    # Truncate sheet name to 31 chars (Excel limit)
    sheet_name = str(entry_no)[-31:]
    ws = wb.create_sheet(title=sheet_name)

    # Title
    ws["A1"] = f"Property Detail: {entry_no}"
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws.merge_cells("A1:D1")
    ws.freeze_panes = "A3"

    row = 2
    # Write header row for detail table
    for header, col in [("Field", "A"), ("Value", "B")]:
        cell = ws[f"{col}{row}"]
        cell.value = header
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN_BORDER
    row += 1

    # Write all property fields
    field_rows = []

    def add_field(name: str, value: Any) -> None:
        nonlocal row
        if isinstance(value, bool):
            value = "Yes" if value else "No"
        elif isinstance(value, (list, dict)):
            value = json.dumps(value, default=str)
        ws[f"A{row}"] = name
        ws[f"A{row}"].font = Font(name="Arial", bold=True, size=9)
        ws[f"A{row}"].border = _THIN_BORDER
        ws[f"B{row}"] = str(value) if value is not None else ""
        ws[f"B{row}"].font = _BODY_FONT
        ws[f"B{row}"].alignment = _WRAP
        ws[f"B{row}"].border = _THIN_BORDER
        row += 1

    # Basic fields
    basic_fields = [
        ("Entry No", "entry_no"), ("Owner", "owner_full_name"), ("Address", "address"),
        ("R Number", "r_number"), ("Instrument No", "instrument_number"),
        ("Sale Date", "sale_date"), ("Legal Description", "legal_description"),
        ("Loan Type", "loan_type"), ("Original Loan Amount", "original_loan_amount"),
        ("Origination Date", "loan_origination_date"), ("Lender", "lender"),
        ("Servicer", "servicer"), ("Trustee", "trustee"),
        ("Is Purchase Money", "is_purchase_money"), ("Has HOA Rider", "has_hoa_rider"),
        ("Status", "status"),
    ]
    for label, key in basic_fields:
        add_field(label, data.get(key))

    # Loan estimates
    add_field("--- Loan Estimate ---", "")
    for label, key in [
        ("Assumed Rate", "_assumed_rate"),
        ("Est Monthly Payment", "_est_monthly_payment"),
        ("Est Current Balance", "_est_balance"),
        ("Est % Paid Down", "_est_pct_paid_down"),
        ("Filter Flag", "_filter_flag"),
    ]:
        val = data.get(key)
        if key == "_est_pct_paid_down" and val is not None:
            val = f"{val * 100:.2f}%"
        elif key == "_assumed_rate" and val is not None:
            val = f"{val * 100:.1f}%"
        add_field(label, val)

    # WCAD
    add_field("County Assessed Value", data.get("_assessed_value"))
    add_field("Lien Count", data.get("_lien_count"))

    # Step results detail
    row += 1
    ws[f"A{row}"] = "Step Results"
    ws[f"A{row}"].font = Font(name="Arial", bold=True, size=10)
    row += 1
    for step_name, step_data in data.get("_step_results", {}).items():
        add_field(f"[{step_name}]", json.dumps(step_data, default=str)[:500])

    # Sources table
    row += 1
    ws[f"A{row}"] = "Sources"
    ws[f"A{row}"].font = Font(name="Arial", bold=True, size=10)
    row += 1
    ws[f"A{row}"] = "Step"
    ws[f"A{row}"].font = _HEADER_FONT
    ws[f"A{row}"].fill = _HEADER_FILL
    ws[f"B{row}"] = "Type"
    ws[f"B{row}"].font = _HEADER_FONT
    ws[f"B{row}"].fill = _HEADER_FILL
    ws[f"C{row}"] = "Link"
    ws[f"C{row}"].font = _HEADER_FONT
    ws[f"C{row}"].fill = _HEADER_FILL
    row += 1

    for src in data.get("_sources", []):
        ws[f"A{row}"] = src.get("step", "")
        ws[f"A{row}"].font = _BODY_FONT
        ws[f"B{row}"] = src.get("type", "")
        ws[f"B{row}"].font = _BODY_FONT
        link_cell = ws[f"C{row}"]
        link_value = src.get("value", "")
        if link_value.startswith("http"):
            link_cell.value = link_value
            link_cell.hyperlink = link_value
            link_cell.font = _LINK_FONT
        else:
            # Local file path → file:// URL
            link_cell.value = link_value
            try:
                link_cell.hyperlink = Path(link_value).as_uri()
                link_cell.font = _LINK_FONT
            except Exception:
                link_cell.font = _BODY_FONT
        row += 1

    # Column widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 80


def run_export(month: str, county: str = "williamson", output_path: str = "out.xlsx") -> None:
    """Generate the Excel export workbook."""
    session = get_session()
    try:
        props = (
            session.query(Property)
            .filter_by(month=month, county=county)
            .order_by(Property.entry_no)
            .all()
        )
        log.info("Exporting %d properties for %s", len(props), month)

        wb = Workbook()
        # Remove default empty sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        # Collect all property data
        properties_data = []
        master_errors = []
        for prop in props:
            try:
                data = _build_property_data(session, prop)
                properties_data.append(data)
            except Exception as exc:
                log.error("Failed to build data for %s: %s", prop.entry_no, exc)
                master_errors.append(prop.entry_no)

        # Write Master sheet first
        try:
            master_ws = wb.create_sheet("Master", 0)
            _write_master_sheet(master_ws, properties_data)
            log.info("Master sheet written: %d rows", len(properties_data))
        except Exception as exc:
            log.error("Master sheet generation failed: %s", exc)
            log.error("Continuing with per-property sheets...")

        # Write per-property tabs
        sheet_errors = []
        for data in properties_data:
            try:
                _write_property_sheet(wb, data)
            except Exception as exc:
                log.error(
                    "Failed to write sheet for %s: %s",
                    data.get("entry_no"), exc,
                )
                sheet_errors.append(data.get("entry_no"))

        # Save workbook
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output))
        log.info("Workbook saved: %s", output)

        # Summary
        if master_errors:
            log.warning(
                "Properties missing from Master sheet (%d): %s",
                len(master_errors), master_errors,
            )
        if sheet_errors:
            log.warning(
                "Properties with failed per-property sheets (%d): %s",
                len(sheet_errors), sheet_errors,
            )
        print(
            f"\n✓ Export complete: {output}\n"
            f"  Properties: {len(properties_data)}\n"
            f"  Master errors: {len(master_errors)}\n"
            f"  Sheet errors: {len(sheet_errors)}"
        )
    except Exception:
        raise
    finally:
        session.close()


def run_csv_export(month: str, county: str = "williamson", output_path: str = "out.csv") -> None:
    """Generate a flat CSV export using the same columns as the Master sheet."""
    session = get_session()
    try:
        props = (
            session.query(Property)
            .filter_by(month=month, county=county)
            .order_by(Property.entry_no)
            .all()
        )
        log.info("CSV export: %d properties for %s", len(props), month)

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        headers = [col[0] for col in _MASTER_COLS]

        errors = 0
        with open(output, "w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for prop in props:
                try:
                    data = _build_property_data(session, prop)
                    row: dict[str, Any] = {}
                    for col_label, key in _MASTER_COLS:
                        value = data.get(key)
                        if key == "_est_pct_paid_down" and value is not None:
                            value = round(value * 100, 2)  # as percent
                        elif isinstance(value, bool):
                            value = "Yes" if value else "No"
                        row[col_label] = value if value is not None else ""
                    writer.writerow(row)
                except Exception as exc:
                    log.error("CSV row failed for %s: %s", prop.entry_no, exc)
                    errors += 1

        log.info("CSV saved: %s", output)
        print(
            f"\n✓ CSV export complete: {output}\n"
            f"  Properties: {len(props) - errors}\n"
            f"  Errors: {errors}"
        )
    except Exception:
        raise
    finally:
        session.close()
